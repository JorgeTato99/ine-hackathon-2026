"""Base de datos para la entrega del Reto 2: agregada por municipio y sin datos personales.

Reúne lo que usó el análisis en una carpeta autocontenida:

- ``historia_municipal.csv``: participación por municipio y elección federal, 2009-2024 (notebook 03).
- ``municipios_2027.csv`` y ``escenarios_2027.csv``: resultados del mapa (notebook 05).
- ``diccionario.csv``: archivo, columna, descripción y unidad de cada columna.
- ``LEEME.md``: qué es, de dónde sale y cómo leerla.

Uso, desde la raíz del repositorio y después de correr los notebooks 03 a 05::

    uv run --locked python -m scripts.entrega_base_datos

Con ``--xlsx`` también escribe un libro de Excel con las mismas tablas (requiere openpyxl, que no es
dependencia del proyecto)::

    uv run --locked --with openpyxl==3.1.5 python -m scripts.entrega_base_datos --xlsx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import polars as pl

RAIZ = Path(__file__).resolve().parents[1]
PROCESADOS = RAIZ / "data" / "processed"
VERSION = "resultados-1.3"
RESULTADOS = PROCESADOS / "mapa" / VERSION
SALIDA = PROCESADOS / "entrega" / "base-datos-1.3"
TEXTO = {"cvegeo": pl.String, "clave_ine_2024": pl.String, "motivo": pl.String, "ultima_eleccion_sin_casilla": pl.String}

HISTORIA = {
    "cvegeo": "clave INEGI del municipio (texto de 5 caracteres, con ceros); territorio de 2024",
    "estado": "nombre de la entidad",
    "municipio": "nombre del municipio",
    "anio": "año de la elección federal (2009, 2015 y 2021 intermedias; 2012, 2018 y 2024 presidenciales)",
    "tipo_eleccion": "presidencial o intermedia",
    "lista_nominal": "personas en la lista nominal de las secciones del municipio (Conteos Censales de Participación del INE)",
    "votaron": "personas que votaron, según los Conteos Censales",
    "no_votaron": "personas que no votaron, según los Conteos Censales",
    "no_sabemos": "personas de las que el registro no dice si votaron",
    "participacion": "porcentaje (0 a 100): votaron / (votaron + no votaron); sale de las actas cuando los conteos no bastan (ver fuente_participacion)",
    "fuente_participacion": "de dónde sale la participación: conteos censales o actas oficiales por casilla (decisión D3)",
    "participacion_nacional": "porcentaje nacional de esa elección, con la misma fórmula",
    "desviacion": "puntos porcentuales: participación del municipio menos la nacional; negativa = votó menos que el país",
    "sin_casilla": "true si al menos 10 % de la lista se quedó sin casilla en esa elección (decisión D12)",
    "pct_sin_casilla": "porcentaje de la lista en secciones sin casilla",
    "territorio_aproximado": "true si el territorio de 2024 no se pudo reconstruir completo ese año (decisión D13)",
    "gobernador_mismo_dia": "true si el estado eligió gubernatura el mismo día",
    "diputacion_local_mismo_dia": "true si el estado eligió congreso local el mismo día",
    "ayuntamiento_mismo_dia": "true si el municipio eligió ayuntamiento el mismo día",
    "usos_y_costumbres": "true si el municipio elegía autoridades por sistemas normativos indígenas en esa elección (Oaxaca y, según el año, hasta 7 municipios fuera de Oaxaca)",
}
COLUMNAS_PANEL = {"AELEC": "anio", "LN": "lista_nominal", "SV": "votaron", "NV": "no_votaron", "NS": "no_sabemos"}
IDENTIDAD = {
    "cvegeo": "clave INEGI del municipio (texto de 5 caracteres, con ceros); clave de unión con la geometría",
    "clave_ine_2024": "clave del INE en 2024 (texto); solo referencia, el INE renumera municipios entre años",
    "nombre_mpio": "nombre del municipio",
    "nombre_edo": "nombre de la entidad",
    "anio": "año de la elección (2027)",
    "nivel_nacional": "escenario nacional: bajo (44.1 %), medio (47.1 %) o alto (51.8 %)",
    "intervencion": "ninguna o dirigida_rojos (supuestos S1 + S2 + S3 en los municipios en rojo; decisión D22)",
    "escenario_id": "nac-<nivel>_int-<intervención>",
    "estado_estimacion": "estimado o sin_estimacion",
    "motivo": "códigos separados por '|': territorio_aproximado, antecedente_sin_casilla",
}


def historia_municipal() -> pl.DataFrame:
    panel = pl.read_parquet(PROCESADOS / "panel_municipal.parquet")
    return (
        panel.rename(COLUMNAS_PANEL)
        .with_columns(usos_y_costumbres=pl.col("es_sni").fill_null(False) | pl.col("sni_fuera_de_oaxaca").fill_null(False))
        .select(list(HISTORIA))
        .sort("cvegeo", "anio")
    )


def diccionario(tablas: dict[str, pl.DataFrame], unidades: dict[str, str]) -> pl.DataFrame:
    def describe(archivo: str, columna: str) -> str:
        if archivo == "historia_municipal.csv":
            return HISTORIA[columna]
        if columna in IDENTIDAD:
            return IDENTIDAD[columna]
        if columna in unidades:
            return unidades[columna]
        if columna in ("part_pred_inf", "part_pred_sup"):
            return unidades["part_pred_inf, part_pred_sup"]
        if columna.startswith("part_"):
            return f"participación observada en {columna[5:]}: " + unidades["part_*"]
        raise KeyError(f"{archivo}: {columna} no tiene descripción")

    filas = [
        {"archivo": archivo, "columna": c, "tipo": str(t.schema[c]), "vacios": t[c].null_count(), "descripcion": describe(archivo, c)}
        for archivo, t in tablas.items()
        for c in t.columns
    ]
    return pl.DataFrame(filas)


LEEME = """# Base de datos · Hackathon INE 2026, Reto 2

Riesgo de abstención por municipio para la elección federal intermedia del 6 de junio de 2027.
**Información agregada por municipio; no contiene datos personales.**

| Archivo | Filas | Contenido |
| --- | ---: | --- |
| `historia_municipal.csv` | {n_hist:,} | Participación de cada municipio en las seis elecciones federales de 2009 a 2024 |
| `municipios_2027.csv` | {n_mun:,} | Semáforo 2027, probabilidad de rojo, personas en juego, alertas y contexto, un municipio por fila |
| `escenarios_2027.csv` | {n_esc:,} | Participación prevista 2027 por municipio, nivel nacional (bajo, medio, alto) y sin o con intervención |
| `diccionario.csv` | {n_dic:,} | Descripción y unidad de cada columna |

**Cómo leerla**

- `cvegeo` es texto de cinco caracteres con ceros a la izquierda (por ejemplo `01001`). Al abrir en Excel,
  importarla como texto.
- Celda vacía = sin dato o no aplica. Nunca equivale a cero.
- Participación = votaron / (votaron + no votaron). En `historia_municipal.csv` va de 0 a 100; en los archivos
  de 2027 va de 0 a 1, como en el mapa.
- El escenario con intervención (`dirigida_rojos`) es una cuenta con supuestos explícitos, **no un efecto
  medido** de una campaña.
- Los niveles nacionales de 2027 (44.06 %, 47.07 % y 51.84 %) son la participación de las intermedias de 2009,
  2015 y 2021 en los Conteos Censales tal como vienen. `participacion_nacional` de la historia usa la
  participación ya corregida con las actas (44.09 %, 47.06 % y 51.85 %). La diferencia es de 0.03 puntos o menos.

**Fuentes:** INE (Conteos Censales de Participación Ciudadana 2009-2024 y cómputos distritales por casilla),
INEGI, CONAPO, CONEVAL, Banxico, SESNSP, IEEPCO y calendarios de los organismos públicos locales.

**Versión:** {version}, generada el {generado}. Las cifras salen de los notebooks 03 a 05 del repositorio.
"""


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--xlsx", action="store_true", help="escribir también un libro de Excel")
    args = parser.parse_args()

    metadatos = json.loads((RESULTADOS / "metadatos.json").read_text(encoding="utf-8"))
    tablas = {
        "historia_municipal.csv": historia_municipal(),
        "municipios_2027.csv": pl.read_csv(RESULTADOS / "municipios.csv", schema_overrides=TEXTO),
        "escenarios_2027.csv": pl.read_csv(RESULTADOS / "escenarios_2027.csv", schema_overrides=TEXTO),
    }
    claves = {archivo: set(t["cvegeo"]) for archivo, t in tablas.items()}
    assert claves["municipios_2027.csv"] == claves["escenarios_2027.csv"], "municipios y escenarios no coinciden"
    assert claves["municipios_2027.csv"] <= claves["historia_municipal.csv"], "hay municipios 2027 sin historia"
    assert all(t["cvegeo"].str.len_chars().eq(5).all() for t in tablas.values()), "cvegeo debe tener 5 caracteres"
    tablas = {a: t.with_columns([pl.col(c).fill_nan(None) for c, tipo in t.schema.items() if tipo.is_float()]) for a, t in tablas.items()}
    dic = diccionario(tablas, metadatos["unidades"])

    SALIDA.mkdir(parents=True, exist_ok=True)
    for archivo, t in tablas.items():
        t.write_csv(SALIDA / archivo)
    dic.write_csv(SALIDA / "diccionario.csv")
    (SALIDA / "LEEME.md").write_text(
        LEEME.format(n_hist=tablas["historia_municipal.csv"].height, n_mun=tablas["municipios_2027.csv"].height,
                     n_esc=tablas["escenarios_2027.csv"].height, n_dic=dic.height,
                     version=metadatos["version_resultados"], generado=metadatos["generado"][:10]),
        encoding="utf-8")
    if args.xlsx:
        from openpyxl import Workbook

        libro = Workbook()
        libro.remove(libro.active)
        for nombre, t in [("diccionario", dic), *((a.removesuffix(".csv"), t) for a, t in tablas.items())]:
            hoja = libro.create_sheet(nombre[:31])
            hoja.append(t.columns)
            for fila in t.iter_rows():
                hoja.append(list(fila))
        libro.save(SALIDA / "base-datos-reto2-ine-2026.xlsx")
    print(f"Base escrita en {SALIDA.relative_to(RAIZ)}:", sorted(p.name for p in SALIDA.iterdir()))


if __name__ == "__main__":
    main()
