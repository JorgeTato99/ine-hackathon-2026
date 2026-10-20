"""CONAPO (proyecciones, marginación, intensidad migratoria) y Banxico (remesas) por municipio.

Fuentes (todas oficiales; la URL exacta y el SHA-256 de cada archivo están en
data/reference/huellas_fuentes_externas.csv y en docs/fuentes-externas.md)
    - CONAPO, proyecciones de población por municipio, base 2020 (1990-2040) y nacional/estatal:
      https://www.datos.gob.mx/dataset/f2b9b220-3ef7-4e3a-bde6-87e1dac78c6a
    - CONAPO, proyecciones municipales base 2015 (2015-2030) y nacional base 2015: copias de la
      Wayback Machine (modo id_, bytes originales) de
      http://www.conapo.gob.mx/work/models/CONAPO/Datos_Abiertos/Proyecciones2018/
    - CONAPO, índice de marginación municipal 2010, 2015 y 2020 (método DP2):
      https://www.datos.gob.mx/dataset/d8f2a534-bcee-4114-853d-82982a81ce24
    - CONAPO, índice de intensidad migratoria 2020 (copia Wayback de repodatos.atdt.gob.mx) y 2010
      (copia Wayback de conapo.gob.mx).
    - Banxico, cuadros CE166 (remesas por municipio, trimestral) y CE81 (nacional mensual) del SIE:
      https://www.banxico.org.mx/SieInternet/

    Este script NO descarga: varios sitios oficiales bloquean descargas automáticas (403 de
    repodatos.atdt.gob.mx, reto JavaScript en gob.mx/conapo, conapo.gob.mx sin respuesta). Los
    crudos se bajan a mano desde las URL de huellas_fuentes_externas.csv a data/raw/externos/
    {conapo,banxico}/ con los nombres que ahí aparecen.

    Herramienta externa: los dos RAR de la base 2015 se extraen a mano con bsdtar (libarchive),
    sin ejecutar nada de su contenido:
        bsdtar -xf data/raw/externos/conapo/base2015_wayback/base_municipios_final_datos_01.rar \
            -C data/raw/externos/conapo/base2015_wayback/extraido
    (igual para _02). El script verifica el SHA-256 de los CSV extraídos.

Qué produce (data/processed/externos/)
    conapo_proyecciones_edad_sexo.parquet           base 2020, 1990-2040, 18 grupos (85+)
    conapo_proyecciones_edad_sexo_base2015.parquet  base 2015, 2015-2030, 14 grupos (65+)
    conapo_marginacion.parquet                      IMM 2010, 2015 y 2020
    conapo_intensidad_migratoria.parquet            IIM 2010 y 2020
    banxico_remesas_municipio.parquet               suma anual 2013-2025
    verificaciones/conapo_banxico_verificaciones.json

Cómo correrlo (desde la raíz del repo; openpyxl y xlrd no están en uv.lock):
    uv tool run --from uv==0.13.1 uv run --no-project --with polars==2.0.0 \
        --with openpyxl==3.1.5 --with xlrd==2.0.1 python -I scripts/externos/fuente_conapo_banxico.py

    Antes de procesar compara el SHA-256 de cada insumo con huellas_fuentes_externas.csv y se
    detiene si difiere (--aceptar-huellas-distintas lo convierte en aviso). El cuadro de Banxico
    lleva la fecha de consulta dentro del archivo, así que una descarga nueva nunca tendrá la
    misma huella: hay que revisar la diferencia y decidir.

Verificaciones (se imprimen y se guardan en el JSON)
    - Proyecciones base 2020: POB_TOTAL = suma de grupos; suma municipal = nacional de CONAPO por
      año y sexo; estructura por edad y sexo = suma de las 32 entidades en años clave.
    - Proyecciones base 2015: suma municipal = nacional de su misma serie.
    - Cruce de claves de cada conjunto contra data/reference/municipios_inegi.csv.
    - Banxico: municipios + "No identificado" = renglón estatal; suma de entidades (CE166) =
      nacional mensual (CE81) por año. El cruce municipal es por nombre exacto normalizado
      dentro de cada entidad, sin coincidencia difusa.
"""

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import polars as pl
import xlrd

RAIZ = next(
    p
    for p in [Path(__file__).resolve(), *Path(__file__).resolve().parents]
    if (p / "pyproject.toml").is_file()
)
HUELLAS = RAIZ / "data" / "reference" / "huellas_fuentes_externas.csv"
RAW = RAIZ / "data/raw/externos"
CON = RAW / "conapo"
BAN = RAW / "banxico"
OUT = RAIZ / "data/processed/externos"
VERIF = OUT / "verificaciones"
CATALOGO = RAIZ / "data/reference/municipios_inegi.csv"

F_PROY2020 = CON / "base2020_datosgob/pobproy_quinq1.csv"
F_IND2020 = CON / "base2020_datosgob/05_indicadores_demograficos_proyecciones.csv"
F_EST2020 = CON / "base2020_wayback/00_Pob_Mitad_1950_2070.csv"
F_PROY2015 = [
    CON / f"base2015_wayback/extraido/base_municipios_final_datos_0{i}.csv" for i in (1, 2)
]
F_NAC2015 = CON / "base2015_wayback/pob_mit_proyecciones.csv"
F_IMM = {
    2010: CON / "marginacion_datosgob/imm_dp2_2010.csv",
    2015: CON / "marginacion_datosgob/imm_dp2_2015.csv",
    2020: CON / "marginacion_datosgob/imm_2020-3.csv",
}
F_IIM2020 = CON / "iim_wayback/06_iim_mex_eeuu_2020_municipio.csv"
F_IIM2010 = CON / "iim2010_wayback/IIM2010_BASEMUN.xls"
F_BANXICO = BAN / "CE166_remesas_municipio_2013T1_2026T2.xlsx"
F_BANXICO_NAC = BAN / "CE81_ingresos_remesas_mensual_2013_2025.xlsx"

ANIOS_CLAVE = [2015, 2018, 2020, 2021, 2024, 2027, 2030]
GRADOS = {
    "muy bajo": "Muy bajo",
    "bajo": "Bajo",
    "medio": "Medio",
    "alto": "Alto",
    "muy alto": "Muy alto",
    "nulo": "Nulo",
}


def sha256(p: Path) -> str:
    with p.open("rb") as f:
        return hashlib.file_digest(f, "sha256").hexdigest()


def huellas_registradas(prefijo: str) -> dict[str, str]:
    """archivo_relativo -> SHA-256 de data/reference/huellas_fuentes_externas.csv bajo `prefijo`."""
    with HUELLAS.open(encoding="utf-8", newline="") as f:
        return {
            r["archivo_relativo"]: r["sha256"]
            for r in csv.DictReader(f)
            if r["archivo_relativo"].startswith(prefijo)
        }


def comprobar_huella(p: Path, esperadas: dict[str, str], estricto: bool) -> str:
    """Devuelve el SHA-256 de `p`; falla (o avisa) si no coincide con la huella registrada."""
    rel = p.relative_to(RAIZ).as_posix()
    h = sha256(p)
    esperado = esperadas.get(rel)
    if esperado == h:
        return h
    problema = "sin huella registrada" if esperado is None else f"SHA-256 {h} != {esperado}"
    if estricto:
        raise SystemExit(f"{rel}: {problema}")
    print(f"AVISO {rel}: {problema}", file=sys.stderr, flush=True)
    return h


def insumos() -> list[Path]:
    """Archivos que lee el script (los PDF de documentación no se leen)."""
    return [
        F_PROY2020,
        F_IND2020,
        F_EST2020,
        *F_PROY2015,
        F_NAC2015,
        *F_IMM.values(),
        F_IIM2020,
        F_IIM2010,
        F_BANXICO,
        F_BANXICO_NAC,
    ]


def leer_csv(p: Path, encoding: str = "utf-8-sig") -> pl.DataFrame:
    texto = p.read_bytes().decode(encoding)
    return pl.read_csv(io.StringIO(texto), infer_schema_length=0)


def cvegeo(col: str) -> pl.Expr:
    return pl.col(col).str.strip_chars().str.replace(r"\.0$", "").str.zfill(5)


def grupo_edad(etiqueta: str) -> str:
    """'POB_010_014' / 'pobm_10_14' -> '10-14'; '..._85_mm' -> '85+'."""
    a, b = etiqueta.split("_")[-2:]
    return f"{int(a)}+" if b == "mm" else f"{int(a)}-{int(b)}"


def catalogo() -> pl.DataFrame:
    return pl.read_csv(
        CATALOGO, schema_overrides={"CVEGEO": pl.String, "CVE_ENT": pl.String, "CVE_MUN": pl.String}
    )


def cruce(nombre: str, claves: set[str], cat: pl.DataFrame) -> dict:
    inegi = set(cat["CVEGEO"].to_list())
    nombres = dict(zip(cat["CVEGEO"], cat["NOMGEO"], strict=False))
    fuera = sorted(claves - inegi)
    faltan = sorted(inegi - claves)
    r = {
        "dataset": nombre,
        "municipios": len(claves),
        "no_en_catalogo_inegi": fuera,
        "catalogo_inegi_sin_dato": [f"{c} {nombres[c]}" for c in faltan],
    }
    print(
        f"[cruce] {nombre}: {len(claves)} claves; fuera del catálogo INEGI {len(fuera)}; INEGI sin dato {len(faltan)}"
    )
    if fuera:
        print("   fuera:", fuera)
    if faltan:
        print("   INEGI sin dato:", r["catalogo_inegi_sin_dato"])
    return r


# ---------------------------------------------------------------- proyecciones


def proyecciones_2020() -> pl.DataFrame:
    d = leer_csv(F_PROY2020)
    cols_edad = [c for c in d.columns if c.startswith("POB_") and c != "POB_TOTAL"]
    d = d.with_columns(
        cvegeo("CLAVE").alias("cvegeo"),
        pl.col("ANO").cast(pl.Int32).alias("anio"),
        pl.col("SEXO")
        .str.to_lowercase()
        .replace({"hombres": "hombre", "mujeres": "mujer"})
        .alias("sexo"),
        *[pl.col(c).cast(pl.Int64) for c in cols_edad + ["POB_TOTAL"]],
    )
    # Verificación interna: POB_TOTAL = suma de grupos quinquenales
    dif = d.select((pl.col("POB_TOTAL") - pl.sum_horizontal(cols_edad)).abs().max()).item()
    print(
        f"[base2020] filas {d.height}; municipios {d['cvegeo'].n_unique()}; años {d['anio'].min()}-{d['anio'].max()}; "
        f"max |POB_TOTAL - suma grupos| = {dif}"
    )
    largo = d.unpivot(
        index=["cvegeo", "anio", "sexo"], on=cols_edad, variable_name="col", value_name="poblacion"
    )
    largo = largo.with_columns(
        pl.col("col").map_elements(grupo_edad, return_dtype=pl.String).alias("grupo_edad")
    )
    return largo.select("cvegeo", "anio", "sexo", "grupo_edad", "poblacion").sort(
        "cvegeo", "anio", "sexo", "grupo_edad"
    )


def proyecciones_2015() -> pl.DataFrame:
    d = pl.concat([leer_csv(p, "latin-1") for p in F_PROY2015])
    d = d.with_columns(
        cvegeo("CLAVE").alias("cvegeo"),
        pl.col("AÑO").cast(pl.Int32).alias("anio"),
        pl.col("SEXO")
        .str.to_lowercase()
        .replace({"hombres": "hombre", "mujeres": "mujer"})
        .alias("sexo"),
        pl.col("EDAD_QUIN").map_elements(grupo_edad, return_dtype=pl.String).alias("grupo_edad"),
        pl.col("POB").cast(pl.Int64).alias("poblacion"),
    )
    print(
        f"[base2015] filas {d.height}; municipios {d['cvegeo'].n_unique()}; años {d['anio'].min()}-{d['anio'].max()}; "
        f"grupos {d['grupo_edad'].unique().sort().to_list()}"
    )
    return d.select("cvegeo", "anio", "sexo", "grupo_edad", "poblacion").sort(
        "cvegeo", "anio", "sexo", "grupo_edad"
    )


def verificar_nacional(p2020: pl.DataFrame, p2015: pl.DataFrame) -> dict:
    res = {}
    # base 2020: nacional publicado (POB_MIT_ANIO, República Mexicana) en indicadores demográficos
    ind = (
        leer_csv(F_IND2020)
        .filter(pl.col("CVE_GEO") == "0")
        .select(
            pl.col("ANIO").cast(pl.Int32).alias("anio"),
            pl.col("POB_MIT_ANIO").cast(pl.Int64).alias("nacional_conapo"),
            pl.col("HOM_MIT_ANIO").cast(pl.Int64).alias("hombres_conapo"),
            pl.col("MUJ_MIT_ANIO").cast(pl.Int64).alias("mujeres_conapo"),
        )
    )
    s = p2020.group_by("anio").agg(
        pl.col("poblacion").sum().alias("suma_municipios"),
        pl.col("poblacion").filter(pl.col("sexo") == "hombre").sum().alias("suma_hombres"),
        pl.col("poblacion").filter(pl.col("sexo") == "mujer").sum().alias("suma_mujeres"),
    )
    v = (
        s.join(ind, on="anio", how="left")
        .with_columns(
            (pl.col("suma_municipios") - pl.col("nacional_conapo")).alias("diferencia"),
            (pl.col("suma_hombres") - pl.col("hombres_conapo")).alias("dif_hombres"),
            (pl.col("suma_mujeres") - pl.col("mujeres_conapo")).alias("dif_mujeres"),
        )
        .sort("anio")
    )
    print(
        "[verif base2020] suma municipal vs POB_MIT_ANIO nacional (CONAPO indicadores 1950-2070):"
    )
    print(v.filter(pl.col("anio").is_in([1990, 2000, 2010] + ANIOS_CLAVE + [2035, 2040])))
    print("   años con diferencia != 0:", v.filter(pl.col("diferencia") != 0)["anio"].to_list())
    res["base2020_vs_nacional"] = v.to_dicts()

    # base 2020: por sexo y grupo quinquenal contra la suma de las 32 entidades (00_Pob_Mitad, edades simples)
    est = (
        leer_csv(F_EST2020)
        .with_columns(
            pl.col("ANIO").cast(pl.Int32).alias("anio"),
            pl.col("EDAD").cast(pl.Int32).alias("edad"),
            pl.col("POBLACION").cast(pl.Int64),
            pl.col("SEXO")
            .str.to_lowercase()
            .replace({"hombres": "hombre", "mujeres": "mujer"})
            .alias("sexo"),
        )
        .with_columns(
            pl.when(pl.col("edad") >= 85)
            .then(pl.lit("85+"))
            .otherwise(
                (pl.col("edad") // 5 * 5).cast(pl.String)
                + "-"
                + (pl.col("edad") // 5 * 5 + 4).cast(pl.String)
            )
            .alias("grupo_edad")
        )
    )
    est = (
        est.filter(pl.col("anio").is_in(ANIOS_CLAVE))
        .group_by("anio", "sexo", "grupo_edad")
        .agg(pl.col("POBLACION").sum().alias("estatal"))
    )
    mun = (
        p2020.filter(pl.col("anio").is_in(ANIOS_CLAVE))
        .group_by("anio", "sexo", "grupo_edad")
        .agg(pl.col("poblacion").sum().alias("municipal"))
    )
    g = mun.join(est, on=["anio", "sexo", "grupo_edad"], how="full", coalesce=True).with_columns(
        (pl.col("municipal") - pl.col("estatal")).alias("dif")
    )
    print(
        f"[verif base2020] celdas año×sexo×grupo comparadas: {g.height}; max |dif| = {g['dif'].abs().max()}; nulos = {g['dif'].null_count()}"
    )
    res["base2020_vs_estatal_edad_sexo_max_abs_dif"] = int(g["dif"].abs().max())

    # base 2015: nacional publicado en pob_mit_proyecciones (República Mexicana, edades simples)
    nac = leer_csv(F_NAC2015, "latin-1")
    nac = nac.filter(pl.col("CVE_GEO") == "0").with_columns(
        pl.col("AÑO").cast(pl.Int32).alias("anio"), pl.col("POBLACION").cast(pl.Int64)
    )
    nac_t = nac.group_by("anio").agg(pl.col("POBLACION").sum().alias("nacional_conapo"))
    s15 = p2015.group_by("anio").agg(pl.col("poblacion").sum().alias("suma_municipios"))
    v15 = (
        s15.join(nac_t, on="anio", how="left")
        .with_columns(
            (pl.col("suma_municipios") - pl.col("nacional_conapo")).alias("diferencia"),
            ((pl.col("suma_municipios") / pl.col("nacional_conapo") - 1) * 100)
            .round(4)
            .alias("dif_pct"),
        )
        .sort("anio")
    )
    print(
        "[verif base2015] suma municipal vs nacional base 2015 (pob_mit_proyecciones, República Mexicana):"
    )
    print(v15)
    res["base2015_vs_nacional"] = v15.to_dicts()

    # comparación entre bases en años clave
    c = (
        p2020.group_by("anio")
        .agg(pl.col("poblacion").sum().alias("base2020"))
        .join(p2015.group_by("anio").agg(pl.col("poblacion").sum().alias("base2015")), on="anio")
        .with_columns(
            ((pl.col("base2015") / pl.col("base2020") - 1) * 100)
            .round(2)
            .alias("base2015_vs_2020_pct")
        )
        .filter(pl.col("anio").is_in(ANIOS_CLAVE))
        .sort("anio")
    )
    print("[comparación] total nacional base 2015 vs base 2020:")
    print(c)
    res["comparacion_bases"] = c.to_dicts()
    return res


# ---------------------------------------------------------------- marginación

MAPA_IMM = {
    "POB_TOT": "poblacion_total",
    "ANALF": "pct_analfabeta_15_mas",
    "SBASC": "pct_sin_educacion_basica_15_mas",
    "OVSDE": "pct_ocupantes_sin_drenaje_ni_sanitario",
    "OVSEE": "pct_ocupantes_sin_energia_electrica",
    "OVSAE": "pct_ocupantes_sin_agua_entubada",
    "OVPT": "pct_ocupantes_piso_tierra",
    "VHAC": "pct_viviendas_hacinamiento",
    "PL_5000": "pct_poblacion_localidades_menos_5000",
    "PO2SM": "pct_ocupados_hasta_2_salarios_minimos",
    "IM": "indice_marginacion",
    "GM": "grado_marginacion",
    "IMN": "indice_marginacion_normalizado",
}


def marginacion() -> pl.DataFrame:
    partes = []
    for anio, p in F_IMM.items():
        d = leer_csv(p)
        ren = {}
        for c in d.columns:
            base = c.replace("PL.5000", "PL_5000").replace(f"_{anio}", "")
            if base in MAPA_IMM:
                ren[c] = MAPA_IMM[base]
        d = d.with_columns(cvegeo("CVE_MUN").alias("cvegeo")).rename(ren)
        num = [v for v in MAPA_IMM.values() if v not in ("grado_marginacion", "poblacion_total")]
        d = d.select(
            "cvegeo",
            pl.lit(anio, pl.Int32).alias("anio"),
            pl.col("poblacion_total").cast(pl.Float64).cast(pl.Int64),
            *[pl.col(c).cast(pl.Float64) for c in num],
            pl.col("grado_marginacion").str.strip_chars().str.to_lowercase().replace(GRADOS),
        )
        print(
            f"[marginación {anio}] {d.height} municipios; grados {d['grado_marginacion'].value_counts(sort=True).rows()}"
        )
        partes.append(d)
    return (
        pl.concat(partes).select("cvegeo", "anio", *list(MAPA_IMM.values())).sort("anio", "cvegeo")
    )


# ---------------------------------------------------------------- intensidad migratoria


def intensidad_migratoria() -> pl.DataFrame:
    d20 = leer_csv(F_IIM2020).select(
        cvegeo("cve_mun").alias("cvegeo"),
        pl.lit(2020, pl.Int32).alias("anio"),
        pl.col("viv_tot").cast(pl.Float64).cast(pl.Int64).alias("total_viviendas"),
        pl.col("viv_rem").cast(pl.Float64).alias("pct_viviendas_reciben_remesas"),
        pl.col("viv_emig").cast(pl.Float64).alias("pct_viviendas_emigrantes_eua"),
        pl.col("viv_circ").cast(pl.Float64).alias("pct_viviendas_migrantes_circulares"),
        pl.col("viv_ret").cast(pl.Float64).alias("pct_viviendas_migrantes_retorno"),
        pl.col("iim_dp2").cast(pl.Float64).alias("indice_intensidad_migratoria"),
        pl.lit(None, pl.Float64).alias("indice_intensidad_migratoria_0a100"),
        pl.col("gim_dp2")
        .str.strip_chars()
        .str.to_lowercase()
        .replace(GRADOS)
        .alias("grado_intensidad_migratoria"),
        pl.col("pos_nal").cast(pl.Float64).cast(pl.Int32).alias("posicion_nacional"),
        pl.lit("DP2").alias("metodo_indice"),
    )
    hoja = xlrd.open_workbook(F_IIM2010).sheet_by_name("IIM2010_BASEMUN")
    enc = hoja.row_values(0)
    filas = [dict(zip(enc, hoja.row_values(i), strict=False)) for i in range(1, hoja.nrows)]
    desc = xlrd.open_workbook(F_IIM2010).sheet_by_name("Descripcion")
    print("[IIM 2010] hoja Descripcion:")
    for i in range(desc.nrows):
        v = [x for x in desc.row_values(i) if str(x).strip()]
        if v:
            print("   ", v)
    d10 = pl.DataFrame(
        {
            "cvegeo": [f"{int(f['ENT']):02d}{int(f['MUN']):03d}" for f in filas],
            "anio": [2010] * len(filas),
            "total_viviendas": [int(f["TOT_VIV"]) for f in filas],
            "pct_viviendas_reciben_remesas": [float(f["VIV_REM"]) for f in filas],
            "pct_viviendas_emigrantes_eua": [float(f["VIV_EMIG"]) for f in filas],
            "pct_viviendas_migrantes_circulares": [float(f["VIV_CIRC"]) for f in filas],
            "pct_viviendas_migrantes_retorno": [float(f["VIV_RET"]) for f in filas],
            "indice_intensidad_migratoria": [float(f["IIM_2010"]) for f in filas],
            "indice_intensidad_migratoria_0a100": [float(f["IIM0a100"]) for f in filas],
            "grado_intensidad_migratoria": [
                GRADOS[re.sub(r"^\d+\s*", "", f["GIM_2010"]).strip().lower()] for f in filas
            ],
            "posicion_nacional": [
                int(f["LUG_NAL"]) if f["LUG_NAL"] != "" else None for f in filas
            ],  # vacío en grado "Nulo"
            "metodo_indice": ["componentes_principales"] * len(filas),
        },
        schema_overrides={"anio": pl.Int32, "posicion_nacional": pl.Int32},
    )
    for d in (d10, d20):
        print(
            f"[IIM {d['anio'][0]}] {d.height} municipios; grados {d['grado_intensidad_migratoria'].value_counts(sort=True).rows()}"
        )
    return pl.concat([d10, d20.select(d10.columns)]).sort("anio", "cvegeo")


# ---------------------------------------------------------------- Banxico

ENT_BANXICO = {
    "Aguascalientes": "01",
    "Baja California": "02",
    "Baja California Sur": "03",
    "Campeche": "04",
    "Coahuila": "05",
    "Colima": "06",
    "Chiapas": "07",
    "Chihuahua": "08",
    "Ciudad de México": "09",
    "Durango": "10",
    "Guanajuato": "11",
    "Guerrero": "12",
    "Hidalgo": "13",
    "Jalisco": "14",
    "México": "15",
    "Michoacán": "16",
    "Morelos": "17",
    "Nayarit": "18",
    "Nuevo León": "19",
    "Oaxaca": "20",
    "Puebla": "21",
    "Querétaro": "22",
    "Quintana Roo": "23",
    "San Luis Potosí": "24",
    "Sinaloa": "25",
    "Sonora": "26",
    "Tabasco": "27",
    "Tamaulipas": "28",
    "Tlaxcala": "29",
    "Veracruz": "30",
    "Yucatán": "31",
    "Zacatecas": "32",
}
TRIM = {"Ene-Mar": 1, "Abr-Jun": 2, "Jul-Sep": 3, "Oct-Dic": 4}


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\bdr\.", "doctor ", s)
    s = re.sub(r"\bgral\.", "general ", s)
    s = re.sub(r"\bnal\.", "nacional ", s)
    s = re.sub(r"\bdistr\.|\bdistrito\b", "dto ", s)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def indice_nombres() -> dict:
    """(cve_ent, nombre normalizado) -> claves, con nombres oficiales de INEGI y de los archivos CONAPO."""
    idx: dict = {}

    def add(code, name, fuente):
        code = str(code).split(".")[0].zfill(5)
        idx.setdefault((code[:2], norm(name)), {}).setdefault(code, set()).add(fuente)

    def filas_csv(p: Path, encoding: str) -> list[dict]:
        with p.open(encoding=encoding, newline="") as f:
            return list(csv.DictReader(f))

    for r in filas_csv(CATALOGO, "utf-8"):
        add(r["CVEGEO"], r["NOMGEO"], "INEGI")
    for anio, p in F_IMM.items():
        for r in filas_csv(p, "utf-8-sig"):
            add(r["CVE_MUN"], r["NOM_MUN"], f"CONAPO IMM {anio}")
    for r in filas_csv(F_IIM2020, "utf-8-sig"):
        add(r["cve_mun"], r["nom_mun"], "CONAPO IIM 2020")
    h = xlrd.open_workbook(F_IIM2010).sheet_by_index(0)
    for i in range(1, h.nrows):
        v = h.row_values(i)
        add(f"{int(v[0]):02d}{int(v[2]):03d}", v[3], "CONAPO IIM 2010")
    for r in filas_csv(F_PROY2020, "utf-8"):
        add(r["CLAVE"], r["NOM_MUN"], "CONAPO proyecciones base 2020")
    return idx


def banxico(cat: pl.DataFrame) -> tuple[pl.DataFrame, dict]:
    ws = openpyxl.load_workbook(F_BANXICO).worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    i_enc = next(i for i, r in enumerate(filas) if r[2] == "Ene-Mar 2013")
    periodos = filas[i_enc][2:]
    print(
        f"[banxico] {filas[3][1]!r} | {filas[6][1]!r} | {filas[1][4]!r} | periodos {periodos[0]} .. {periodos[-1]} ({len(periodos)})"
    )
    notas = [
        r[1] for r in filas if r[1] and (r[1].startswith("Notas") or r[1].startswith("Las cifras"))
    ]
    print("[banxico] notas del cuadro:", notas)
    idx = indice_nombres()
    registros, sin_cruce, no_ident, estatal = [], [], [], {}
    ent = None
    for r in filas[i_enc + 1 :]:
        etiqueta = r[1]
        if not etiqueta:
            continue
        if etiqueta.startswith(" ●"):
            nom_ent = etiqueta.strip(" ●  ")
            ent = ENT_BANXICO[nom_ent]
            estatal[ent] = r[2:]
            continue
        if not etiqueta.startswith("  ⚬"):
            continue
        nom = etiqueta.strip("  ⚬ ")
        if nom == "No identificado":
            no_ident.append((ent, r[2:]))
            continue
        cand = idx.get((ent, norm(nom)), {})
        if len(cand) != 1:
            sin_cruce.append({"cve_ent": ent, "nombre_banxico": nom, "candidatos": sorted(cand)})
            clave = None
        else:
            clave = next(iter(cand))
        for per, val in zip(periodos, r[2:], strict=False):
            t, a = per.split(" ")
            registros.append(
                (ent, nom, clave, int(a), TRIM[t], None if val is None or val == "" else float(val))
            )
    q = pl.DataFrame(
        registros,
        schema=["cve_ent", "nombre_banxico", "cvegeo", "anio", "trimestre", "remesas_millones_usd"],
        orient="row",
    )
    print(
        f"[banxico] municipios en el cuadro: {q.select('cve_ent', 'nombre_banxico').unique().height}; "
        f"sin cruce: {len(sin_cruce)} -> {sin_cruce}"
    )

    # Verificación interna: municipios + 'No identificado' = renglón estatal
    ver = []
    for ent, vals in no_ident:
        s = (
            q.filter(pl.col("cve_ent") == ent)
            .group_by("anio", "trimestre")
            .agg(pl.col("remesas_millones_usd").sum())
            .sort("anio", "trimestre")
        )
        tot = [
            a + (b or 0) for a, b in zip(s["remesas_millones_usd"].to_list(), vals, strict=False)
        ]
        dif = max(abs(x - (y or 0)) for x, y in zip(tot, estatal[ent], strict=False))
        ver.append(dif)
    print(
        f"[banxico] max |suma municipios + no identificado - renglón estatal| = {max(ver):.6f} millones USD"
    )
    tot_ni = sum(sum(v or 0 for v in vals[: len(periodos) - 2]) for _, vals in no_ident)
    tot_est = sum(sum(v or 0 for v in vals[: len(periodos) - 2]) for vals in estatal.values())
    print(
        f"[banxico] 2013T1-2025T4: total estatal {tot_est:,.1f} M USD; 'No identificado' {tot_ni:,.1f} ({tot_ni / tot_est * 100:.2f}%)"
    )
    est_anual = {}
    for vals in estatal.values():
        for per, v in zip(periodos, vals, strict=False):
            a = int(per.split(" ")[1])
            est_anual[a] = est_anual.get(a, 0) + (v or 0)
    print(
        "[banxico] suma anual de las 32 entidades (M USD):",
        {a: round(v, 1) for a, v in est_anual.items()},
    )
    # Contraste con el total nacional mensual publicado (CE81, 'Remesas Totales (Millones de dólares)')
    f81 = list(openpyxl.load_workbook(F_BANXICO_NAC).worksheets[0].iter_rows(values_only=True))
    meses = next(r for r in f81 if r[2] == "Ene 2013")[2:]
    tot = next(r for r in f81 if r[1] and "Remesas Totales (Millones" in r[1])[2:]
    nac81 = {}
    for m, v in zip(meses, tot, strict=False):
        a = int(m.split(" ")[1])
        nac81[a] = nac81.get(a, 0) + v
    comp = {
        a: (round(est_anual[a], 1), round(nac81[a], 1), round(est_anual[a] - nac81[a], 3))
        for a in sorted(nac81)
    }
    print("[banxico] año: (suma 32 entidades CE166, nacional CE81, diferencia) M USD:", comp)

    anual = (
        q.filter(pl.col("cvegeo").is_not_null() & (pl.col("anio") <= 2025))
        .group_by("cvegeo", "anio")
        .agg(
            pl.col("remesas_millones_usd").sum(),
            pl.col("remesas_millones_usd")
            .is_not_null()
            .sum()
            .cast(pl.Int32)
            .alias("trimestres_con_dato"),
            (pl.col("remesas_millones_usd") > 0)
            .sum()
            .cast(pl.Int32)
            .alias("trimestres_con_monto_positivo"),
        )
        .with_columns(pl.col("anio").cast(pl.Int32))
        .sort("cvegeo", "anio")
    )
    info = {
        "sin_cruce": sin_cruce,
        "no_identificado_total_2013_2025": tot_ni,
        "total_estatal_2013_2025": tot_est,
        "suma_anual_entidades": est_anual,
        "entidades_vs_nacional_CE81": comp,
    }
    ceros = anual.filter(pl.col("trimestres_con_monto_positivo") == 0)
    print(
        f"[banxico] municipio-año con 4 trimestres en cero: {ceros.height} (municipios distintos {ceros['cvegeo'].n_unique()})"
    )
    return anual, info


# ---------------------------------------------------------------- main


def main() -> None:
    ap = argparse.ArgumentParser(description="CONAPO y Banxico por municipio.")
    ap.add_argument(
        "--aceptar-huellas-distintas",
        action="store_true",
        help="avisa en vez de fallar si un insumo no coincide con huellas_fuentes_externas.csv",
    )
    a = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    VERIF.mkdir(parents=True, exist_ok=True)
    print("== SHA-256 de los insumos (contra huellas_fuentes_externas.csv)")
    esperadas = huellas_registradas("data/raw/externos/")
    huellas = {}
    for p in insumos():
        h = comprobar_huella(p, esperadas, not a.aceptar_huellas_distintas)
        huellas[p.relative_to(RAIZ).as_posix()] = h
        print(f"{h}  {p.relative_to(RAW)}  ({p.stat().st_size:,} bytes)")

    cat = catalogo()
    print(f"\n== Catálogo INEGI: {cat.height} municipios")
    cruces = []

    print("\n== Proyecciones")
    p20 = proyecciones_2020()
    p15 = proyecciones_2015()
    ver = verificar_nacional(p20, p15)
    p20.write_parquet(OUT / "conapo_proyecciones_edad_sexo.parquet")
    p15.write_parquet(OUT / "conapo_proyecciones_edad_sexo_base2015.parquet")
    cruces.append(cruce("proyecciones base 2020", set(p20["cvegeo"]), cat))
    cruces.append(cruce("proyecciones base 2015", set(p15["cvegeo"]), cat))
    nuevos = sorted(set(p20["cvegeo"]) - set(p15["cvegeo"]))
    nombres = dict(zip(cat["CVEGEO"], cat["NOMGEO"], strict=False))
    print(
        "[proyecciones] en base 2020 pero no en base 2015:",
        [f"{c} {nombres.get(c)}" for c in nuevos],
    )

    print("\n== Marginación")
    m = marginacion()
    m.write_parquet(OUT / "conapo_marginacion.parquet")
    for a in (2010, 2015, 2020):
        cruces.append(cruce(f"marginación {a}", set(m.filter(pl.col("anio") == a)["cvegeo"]), cat))

    print("\n== Intensidad migratoria")
    iim = intensidad_migratoria()
    iim.write_parquet(OUT / "conapo_intensidad_migratoria.parquet")
    for a in (2010, 2020):
        cruces.append(cruce(f"IIM {a}", set(iim.filter(pl.col("anio") == a)["cvegeo"]), cat))
    # Contraste con el libro IIM 2020 (Anexo B, Cuadro B.31): Suma, Yucatán
    print(
        "[IIM 2020] 31072 (Suma):",
        iim.filter((pl.col("cvegeo") == "31072") & (pl.col("anio") == 2020)).to_dicts(),
    )

    print("\n== Banxico")
    b, binfo = banxico(cat)
    b.write_parquet(OUT / "banxico_remesas_municipio.parquet")
    cruces.append(cruce("Banxico remesas", set(b["cvegeo"]), cat))

    resumen = {
        "huellas_insumos": huellas,
        "verificaciones_proyecciones": ver,
        "cruces": cruces,
        "banxico": binfo,
    }
    (VERIF / "conapo_banxico_verificaciones.json").write_text(
        json.dumps(resumen, ensure_ascii=False, indent=1, default=str), encoding="utf-8"
    )
    print("\n== Parquets")
    for f in sorted(OUT.glob("conapo_*.parquet")) + sorted(OUT.glob("banxico_*.parquet")):
        d = pl.read_parquet(f)
        print(f"{f.name}: {d.shape} {d.schema}")


if __name__ == "__main__":
    main()
