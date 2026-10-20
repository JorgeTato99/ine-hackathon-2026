"""CONEVAL municipal: pobreza 2010-2015-2020 (serie comparable) e Índice de Rezago Social 2000-2020.

Fuente
    CONEVAL, https://www.coneval.org.mx
    - Pobreza municipal 2010-2020, "Anexo estadístico" (página
      https://www.coneval.org.mx/Medicion/Paginas/Pobreza-municipio-2010-2020.aspx):
      https://www.coneval.org.mx/Medicion/Documents/Pobreza_municipal/2020/Concentrado_indicadores_de_pobreza_2020.zip
    - Índice de Rezago Social 2000-2020, entidades y municipios (página
      https://www.coneval.org.mx/Medicion/IRS/Paginas/Indice_de_Rezago_Social_2020_anexos.aspx):
      https://www.coneval.org.mx/Medicion/Documents/IRS_2020/IRS_ent_mun_2000_2020.zip

Qué produce (data/processed/externos/)
    coneval_pobreza_municipal.parquet   una fila por municipio y año (2010, 2015, 2020)
    coneval_rezago_social.parquet       una fila por municipio y año (2000 a 2020, cada 5 años)
    verificaciones/coneval_verificaciones.json

    Además extrae los XLSX de los ZIP en data/raw/externos/coneval/extraido/.

Cómo correrlo (desde la raíz del repo; openpyxl no está en uv.lock):
    uv tool run --from uv==0.13.1 uv run --no-project --with polars==2.0.0 \
        --with openpyxl==3.1.5 python -I scripts/externos/fuente_coneval.py [--descargar]

    Sin --descargar usa los ZIP ya presentes en data/raw/externos/coneval/. Con --descargar los
    vuelve a bajar. En ambos casos el SHA-256 de cada ZIP y de cada XLSX extraído debe coincidir
    con data/reference/huellas_fuentes_externas.csv o el script se detiene.

Verificaciones (en el JSON)
    - Encabezados esperados en cada hoja; columnas sin encabezado deben venir vacías.
    - Pobreza: sin duplicados cvegeo-año; porcentajes en [0, 100]; pobreza = extrema + moderada y
      los cuatro grupos suman 100 (con tolerancia de redondeo); suma municipal de personas =
      hoja estatal del mismo archivo; cifras del Comunicado 016/2021 (municipios que concentran
      la mitad de la población en pobreza, cambios de 5 puntos, 15 municipios extremos).
    - Rezago social: municipios y grados por año; porcentajes en [0, 100]; media y desviación
      del índice por año.
    - Cruce de claves contra data/reference/municipios_inegi.csv.

Seguridad
    Solo descarga por HTTPS desde www.coneval.org.mx. Del ZIP solo se aceptan XLSX en la raíz
    (sin rutas, para evitar path traversal); se leen como datos con openpyxl.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
import urllib.request
import zipfile
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl
import polars as pl

RAIZ = next(
    p
    for p in [Path(__file__).resolve(), *Path(__file__).resolve().parents]
    if (p / "pyproject.toml").is_file()
)
HUELLAS = RAIZ / "data" / "reference" / "huellas_fuentes_externas.csv"
RAW = RAIZ / "data/raw/externos/coneval"
OUT = RAIZ / "data/processed/externos"
REF = RAIZ / "data/reference/municipios_inegi.csv"
VERIF_JSON = OUT / "verificaciones" / "coneval_verificaciones.json"

BASE = "https://www.coneval.org.mx"
# Nombre local -> ruta en BASE. El SHA-256 esperado está en huellas_fuentes_externas.csv.
FUENTES = {
    "Concentrado_indicadores_de_pobreza_2020.zip": "/Medicion/Documents/Pobreza_municipal/2020/Concentrado_indicadores_de_pobreza_2020.zip",
    "IRS_ent_mun_2000_2020.zip": "/Medicion/Documents/IRS_2020/IRS_ent_mun_2000_2020.zip",
}

# ---------------------------------------------------------------- utilidades


def sha256(p: Path) -> str:
    with p.open("rb") as f:
        return hashlib.file_digest(f, "sha256").hexdigest()


def huellas_registradas(prefijo: str) -> dict[str, str]:
    """archivo_relativo -> SHA-256 de data/reference/huellas_fuentes_externas.csv bajo `prefijo`."""
    with HUELLAS.open(encoding="utf-8", newline="") as f:
        return {
            r["archivo_relativo"]: r["sha256"]
            for r in csv.DictReader(f)
            if r["archivo_relativo"].startswith(prefijo)
        }


def comprobar_huella(p: Path, esperadas: dict[str, str]) -> str:
    """Devuelve el SHA-256 de `p`; se detiene si no coincide con la huella registrada."""
    rel = p.relative_to(RAIZ).as_posix()
    h = sha256(p)
    if esperadas.get(rel) != h:
        raise SystemExit(f"{rel}: SHA-256 {h} != {esperadas.get(rel, 'sin huella registrada')}")
    return h


def norm(s: object) -> str:
    """Texto de encabezado normalizado: sin acentos, minúsculas, espacios simples."""
    if s is None:
        return ""
    t = unicodedata.normalize("NFKD", str(s))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().lower()


def exigir(condicion: object, detalle: object) -> None:
    """Se detiene si la estructura del archivo no es la esperada (no depende de `python -O`)."""
    if not condicion:
        raise SystemExit(f"estructura inesperada: {detalle!r}")


def a_num(v: object) -> float | None:
    """Número o None. 'n.d.' / 'n.a.' / '' -> None. Cualquier otro texto es error."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().lower()
    if s in {"", "n.d.", "n.a.", "nd", "na", "-"}:
        return None
    raise ValueError(f"valor no numérico inesperado: {v!r}")


def descargar() -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    for nombre, ruta in FUENTES.items():
        url = BASE + ruta
        partes = urlsplit(url)
        if partes.scheme != "https" or partes.hostname != "www.coneval.org.mx":
            raise ValueError(f"URL fuera de la fuente oficial: {url}")
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        tmp = RAW / (nombre + ".part")
        # Esquema y host validados arriba.
        with urllib.request.urlopen(req, timeout=300) as r, tmp.open("wb") as f:  # nosec B310
            f.write(r.read())
        tmp.replace(RAW / nombre)


def verificar_y_extraer() -> dict[str, str]:
    esperadas = huellas_registradas("data/raw/externos/coneval/")
    hashes = {}
    for nombre in FUENTES:
        hashes[nombre] = comprobar_huella(RAW / nombre, esperadas)
    destinos = {
        "Concentrado_indicadores_de_pobreza_2020.zip": RAW / "extraido/pobreza",
        "IRS_ent_mun_2000_2020.zip": RAW / "extraido/irs",
    }
    for nombre, dest in destinos.items():
        dest.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(RAW / nombre) as z:
            for m in z.infolist():
                # Solo XLSX planos, sin rutas (evita path traversal).
                if (
                    m.is_dir()
                    or "/" in m.filename
                    or "\\" in m.filename
                    or not m.filename.endswith(".xlsx")
                ):
                    raise SystemExit(f"miembro inesperado en {nombre}: {m.filename}")
                (dest / m.filename).write_bytes(z.read(m))
                hashes[f"extraido/{dest.name}/{m.filename}"] = comprobar_huella(
                    dest / m.filename, esperadas
                )
    return hashes


# ---------------------------------------------------------------- pobreza

INDICADORES_POBREZA = {
    "pobreza": "pobreza",
    "pobreza extrema": "pobreza_extrema",
    "pobreza moderada": "pobreza_moderada",
    "vulnerables por carencia social": "vulnerable_carencias",
    "vulnerables por ingreso": "vulnerable_ingreso",
    "no pobre y no vulnerable": "no_pobre_no_vulnerable",
    "rezago educativo": "carencia_rezago_educativo",
    "carencia por acceso a los servicios de salud": "carencia_salud",
    "carencia por acceso a la seguridad social": "carencia_seguridad_social",
    "carencia por calidad y espacios de la vivienda": "carencia_calidad_vivienda",
    "carencia por acceso a los servicios basicos en la vivienda": "carencia_servicios_basicos_vivienda",
    "carencia por acceso a la alimentacion": "carencia_alimentacion",
    "poblacion con al menos una carencia social": "al_menos_una_carencia",
    "poblacion con tres o mas carencias sociales": "tres_o_mas_carencias",
    "poblacion con ingreso inferior a la linea de pobreza por ingresos": "ingreso_inferior_lpi",
    "poblacion con ingreso inferior a la linea de pobreza extrema por ingresos": "ingreso_inferior_lpei",
}
MEDIDAS = {"porcentaje": "pct", "personas": "personas", "carencias promedio": "carencias_promedio"}
ANIOS_POB = (2010, 2015, 2020)


def mapa_columnas_pobreza(h1: tuple, h2: tuple, col0: int) -> dict[int, tuple[str, int]]:
    """Columna -> (nombre_salida, año) a partir de las dos filas de encabezado."""
    mapa: dict[int, tuple[str, int]] = {}
    grupo = None
    for j in range(col0, len(h1)):
        if h1[j] is not None:
            grupo = norm(h1[j])
        sub = norm(h2[j]) if j < len(h2) else ""
        if not sub:
            continue
        m = re.fullmatch(r"(porcentaje|personas|carencias promedio) (2010|2015|2020)", sub)
        if not m or grupo not in INDICADORES_POBREZA:
            raise SystemExit(f"encabezado no reconocido col {j}: {grupo!r} / {sub!r}")
        mapa[j] = (f"{INDICADORES_POBREZA[grupo]}_{MEDIDAS[m.group(1)]}", int(m.group(2)))
    return mapa


def leer_pobreza(xlsx: Path) -> tuple[pl.DataFrame, pl.DataFrame, dict]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    info: dict = {}

    # --- municipal
    rows = list(wb["Concentrado municipal"].iter_rows(values_only=True))
    h1, h2 = rows[4], rows[5]
    exigir(
        [norm(x) for x in h1[1:5]]
        == ["clave de entidad", "entidad federativa", "clave de municipio", "municipio"],
        h1[1:5],
    )
    pob_cols = {}
    for j in (5, 6, 7):
        m = re.match(r"poblacion (2010|2015|2020)", norm(h1[j]))
        exigir(m, h1[j])
        pob_cols[j] = int(m.group(1))
    mapa = mapa_columnas_pobreza(h1, h2, 8)
    # La última columna (146) solo hereda el encabezado combinado: debe venir vacía.
    extra = [j for j in range(8, len(h1)) if j not in mapa]
    registros = []
    nd_na = 0
    for r in rows[6:]:
        clave = r[3]
        if not (isinstance(clave, str) and re.fullmatch(r"\d{5}", clave)):
            continue
        for j in extra:
            if r[j] is not None:
                raise SystemExit(f"columna sin encabezado con datos: {j} {r[j]!r}")
        exigir(clave[:2] == r[1], (clave, r[1]))
        for anio in ANIOS_POB:
            reg = {
                "cvegeo": clave,
                "anio": anio,
                "cve_ent": r[1],
                "entidad": r[2],
                "municipio": r[4],
            }
            for j, a in pob_cols.items():
                if a == anio:
                    reg["poblacion"] = a_num(r[j])
            for j, (col, a) in mapa.items():
                if a == anio:
                    if isinstance(r[j], str):
                        nd_na += 1
                    reg[col] = a_num(r[j])
            registros.append(reg)
    info["celdas_texto_nd_na_municipal"] = nd_na
    mun = pl.DataFrame(registros, infer_schema_length=None)

    # --- estatal (solo para verificación)
    rows_e = list(wb["Concentrado estatal"].iter_rows(values_only=True))
    e1, e2 = rows_e[4], rows_e[5]
    mapa_e = mapa_columnas_pobreza(e1, e2, 6)
    reg_e = []
    for r in rows_e[6:]:
        if not (isinstance(r[1], str) and re.fullmatch(r"\d{2}", r[1])):
            continue
        for anio in ANIOS_POB:
            reg = {"cve_ent": r[1], "anio": anio}
            for j, (col, a) in mapa_e.items():
                if a == anio:
                    reg[col] = a_num(r[j])
            reg_e.append(reg)
    est = pl.DataFrame(reg_e, infer_schema_length=None)
    return mun, est, info


# ---------------------------------------------------------------- rezago social

IRS_ENCABEZADOS = [
    ("poblacion de 15 anos o mas analfabeta", "pct_analfabeta_15_mas"),
    ("poblacion de 6 a 14 anos que no asiste a la escuela", "pct_no_asiste_escuela_6_14"),
    (
        "poblacion de 15 anos y mas con educacion basica incompleta",
        "pct_educacion_basica_incompleta_15_mas",
    ),
    ("poblacion sin derechohabiencia a servicios de salud", "pct_sin_derechohabiencia_salud"),
    ("viviendas con piso de tierra", "pct_viviendas_piso_tierra"),
    ("viviendas que no disponen de excusado o sanitario", "pct_viviendas_sin_excusado"),
    (
        "viviendas que no disponen de agua entubada de la red publica",
        "pct_viviendas_sin_agua_entubada",
    ),
    ("viviendas que no disponen de drenaje", "pct_viviendas_sin_drenaje"),
    ("viviendas que no disponen de energia electrica", "pct_viviendas_sin_electricidad"),
    ("viviendas que no disponen de lavadora", "pct_viviendas_sin_lavadora"),
    ("viviendas que no disponen de refrigerador", "pct_viviendas_sin_refrigerador"),
]
GRADOS = ["Muy bajo", "Bajo", "Medio", "Alto", "Muy alto"]


def leer_irs(xlsx: Path, anio: int) -> pl.DataFrame:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["Municipios"].iter_rows(values_only=True))
    h1, h2 = rows[4], rows[5]
    esperado1 = [
        "clave entidad",
        "entidad federativa",
        "clave municipio",
        "municipio",
        "poblacion total",
    ]
    exigir([norm(x) for x in h1[:5]] == esperado1, h1[:5])
    exigir([norm(x) for x in h2[5:16]] == [a for a, _ in IRS_ENCABEZADOS], h2[5:16])
    exigir(
        [norm(x) for x in h1[16:19]]
        == [
            "indice de rezago social",
            "grado de rezago social",
            "lugar que ocupa en el contexto nacional",
        ],
        h1[16:19],
    )
    exigir(str(anio) in str(rows[0][2]), rows[0])
    regs = []
    for r in rows[6:]:
        clave = r[2]
        if not (isinstance(clave, str) and re.fullmatch(r"\d{5}", clave)):
            continue
        exigir(clave[:2] == r[0], (clave, r[0]))
        grado = (r[17] or "").strip() or None
        if grado is not None and grado not in GRADOS:
            raise SystemExit(f"grado inesperado {grado!r} en {clave} {anio}")
        reg = {
            "cvegeo": clave,
            "anio": anio,
            "cve_ent": r[0],
            "entidad": r[1],
            "municipio": r[3],
            "poblacion_total": a_num(r[4]),
        }
        for k, (_, col) in enumerate(IRS_ENCABEZADOS):
            reg[col] = a_num(r[5 + k])
        reg["indice_rezago_social"] = a_num(r[16])
        reg["grado_rezago_social"] = grado
        reg["lugar_nacional"] = a_num(r[18])
        regs.append(reg)
    return pl.DataFrame(regs, infer_schema_length=None)


# ---------------------------------------------------------------- verificaciones


def municipios_para_mitad(df: pl.DataFrame, col: str) -> int:
    s = df.select(pl.col(col).drop_nulls().sort(descending=True)).to_series()
    total = s.sum()
    return int((s.cum_sum() < total / 2).sum() + 1)


def verificar(pob: pl.DataFrame, est: pl.DataFrame, irs: pl.DataFrame) -> dict:
    v: dict = {}
    ref = pl.read_csv(
        REF, schema_overrides={"CVEGEO": pl.String, "CVE_ENT": pl.String, "CVE_MUN": pl.String}
    )
    ref_set = set(ref["CVEGEO"].to_list())
    v["municipios_referencia"] = len(ref_set)

    # --- pobreza
    vp: dict = {}
    vp["filas"] = pob.height
    vp["claves_unicas"] = pob["cvegeo"].n_unique()
    vp["duplicados_cvegeo_anio"] = pob.height - pob.select("cvegeo", "anio").unique().height
    vp["municipios_con_estimacion_pobreza"] = {
        str(a): int(pob.filter((pl.col("anio") == a) & pl.col("pobreza_pct").is_not_null()).height)
        for a in ANIOS_POB
    }
    sin = (
        pob.filter(pl.col("pobreza_pct").is_null())
        .select("anio", "cvegeo", "municipio")
        .sort("anio", "cvegeo")
    )
    vp["sin_estimacion"] = {
        str(a): [f"{c} {m}" for c, m in zip(g["cvegeo"], g["municipio"], strict=False)]
        for (a,), g in sin.group_by("anio", maintain_order=True)
    }
    pct_cols = [c for c in pob.columns if c.endswith("_pct")]
    vp["rango_pct"] = {
        "min": float(pob.select(pl.min_horizontal(pl.col(pct_cols).min())).item()),
        "max": float(pob.select(pl.max_horizontal(pl.col(pct_cols).max())).item()),
    }
    vp["pct_fuera_0_100"] = int(
        sum(pob.filter((pl.col(c) < 0) | (pl.col(c) > 100)).height for c in pct_cols)
    )
    # Identidades: pobreza = extrema + moderada; 4 grupos suman 100
    ok = pob.filter(pl.col("pobreza_pct").is_not_null())
    vp["max_abs_pobreza_menos_ext_mod_pct"] = float(
        ok.select(
            (pl.col("pobreza_pct") - pl.col("pobreza_extrema_pct") - pl.col("pobreza_moderada_pct"))
            .abs()
            .max()
        ).item()
    )
    vp["max_abs_cuatro_grupos_menos_100"] = float(
        ok.select(
            (
                pl.col("pobreza_pct")
                + pl.col("vulnerable_carencias_pct")
                + pl.col("vulnerable_ingreso_pct")
                + pl.col("no_pobre_no_vulnerable_pct")
                - 100
            )
            .abs()
            .max()
        ).item()
    )
    # Suma municipal vs hoja estatal (personas)
    pers = [c for c in pob.columns if c.endswith("_personas")]
    suma = pob.group_by("cve_ent", "anio").agg([pl.col(c).sum() for c in pers + ["poblacion"]])
    j = suma.join(est, on=["cve_ent", "anio"], suffix="_est")
    difs = {c: float(j.select((pl.col(c) - pl.col(f"{c}_est")).abs().max()).item()) for c in pers}
    vp["max_abs_dif_personas_suma_municipal_vs_hoja_estatal"] = difs
    nac = (
        pob.group_by("anio")
        .agg(
            pl.col("pobreza_personas").sum(),
            pl.col("poblacion").sum(),
            pl.col("pobreza_extrema_personas").sum(),
        )
        .sort("anio")
    )
    vp["nacional_desde_suma_municipal"] = {
        str(r["anio"]): {
            "poblacion": r["poblacion"],
            "pobreza_personas": r["pobreza_personas"],
            "pobreza_pct": round(100 * r["pobreza_personas"] / r["poblacion"], 2),
            "pobreza_extrema_personas": r["pobreza_extrema_personas"],
        }
        for r in nac.iter_rows(named=True)
    }
    # Contra el comunicado 016/2021 de CONEVAL
    vp["mitad_pobres_en_n_municipios"] = {
        str(a): municipios_para_mitad(pob.filter(pl.col("anio") == a), "pobreza_personas")
        for a in ANIOS_POB
    }
    vp["mitad_pobres_extremos_en_n_municipios_2020"] = municipios_para_mitad(
        pob.filter(pl.col("anio") == 2020), "pobreza_extrema_personas"
    )
    vp["mitad_no_pobres_no_vulnerables_en_n_municipios_2020"] = municipios_para_mitad(
        pob.filter(pl.col("anio") == 2020), "no_pobre_no_vulnerable_personas"
    )
    w = pob.pivot(on="anio", index="cvegeo", values="pobreza_pct")
    d = w.select((pl.col("2020") - pl.col("2015")).alias("d")).drop_nulls()["d"]
    vp["n_baja_5pp_o_mas_2015_2020"] = int((d <= -5).sum())
    vp["n_sube_5pp_o_mas_2015_2020"] = int((d >= 5).sum())
    t15 = (
        pob.filter(pl.col("anio") == 2020)
        .sort("pobreza_pct", descending=True, nulls_last=True)
        .head(15)
    )
    vp["top15_pobreza_pct_2020"] = [
        f"{c} {m} {p:.1f}"
        for c, m, p in zip(t15["cvegeo"], t15["municipio"], t15["pobreza_pct"], strict=False)
    ]
    b15 = (
        pob.filter((pl.col("anio") == 2020) & pl.col("pobreza_pct").is_not_null())
        .sort("pobreza_pct")
        .head(15)
    )
    vp["bottom15_pobreza_pct_2020"] = [
        f"{c} {m} {p:.1f}"
        for c, m, p in zip(b15["cvegeo"], b15["municipio"], b15["pobreza_pct"], strict=False)
    ]
    claves = set(pob["cvegeo"].to_list())
    vp["claves_coneval_no_en_referencia"] = sorted(claves - ref_set)
    vp["claves_referencia_no_en_coneval"] = sorted(
        f"{c} {n}" for c, n in zip(ref["CVEGEO"], ref["NOMGEO"], strict=False) if c not in claves
    )
    v["pobreza"] = vp

    # --- rezago social
    vi: dict = {}
    vi["filas"] = irs.height
    vi["duplicados_cvegeo_anio"] = irs.height - irs.select("cvegeo", "anio").unique().height
    vi["municipios_por_anio"] = {
        str(a): int(n) for a, n in irs.group_by("anio").len().sort("anio").iter_rows()
    }
    vi["con_indice_por_anio"] = {
        str(a): int(n)
        for a, n in irs.filter(pl.col("indice_rezago_social").is_not_null())
        .group_by("anio")
        .len()
        .sort("anio")
        .iter_rows()
    }
    vi["grados_por_anio"] = {
        str(a): {
            g: int(irs.filter((pl.col("anio") == a) & (pl.col("grado_rezago_social") == g)).height)
            for g in GRADOS
        }
        for a in (2000, 2005, 2010, 2015, 2020)
    }
    pc = [c for c in irs.columns if c.startswith("pct_")]
    vi["rango_pct"] = {
        "min": float(irs.select(pl.min_horizontal(pl.col(pc).min())).item()),
        "max": float(irs.select(pl.max_horizontal(pl.col(pc).max())).item()),
    }
    vi["pct_fuera_0_100"] = int(
        sum(irs.filter((pl.col(c) < 0) | (pl.col(c) > 100)).height for c in pc)
    )
    vi["indice_media_y_sd_por_anio"] = {
        str(r["anio"]): [round(r["m"], 6), round(r["s"], 6)]
        for r in irs.group_by("anio")
        .agg(
            pl.col("indice_rezago_social").mean().alias("m"),
            pl.col("indice_rezago_social").std().alias("s"),
        )
        .sort("anio")
        .iter_rows(named=True)
    }
    t10 = (
        irs.filter(pl.col("anio") == 2020)
        .sort("indice_rezago_social", descending=True, nulls_last=True)
        .head(10)
    )
    vi["top10_irs_2020"] = [
        f"{c} {m} {x:.3f}"
        for c, m, x in zip(
            t10["cvegeo"], t10["municipio"], t10["indice_rezago_social"], strict=False
        )
    ]
    vi["claves_no_en_referencia_por_anio"] = {
        str(a): sorted(set(irs.filter(pl.col("anio") == a)["cvegeo"].to_list()) - ref_set)
        for a in (2000, 2005, 2010, 2015, 2020)
    }
    vi["referencia_no_en_irs_por_anio"] = {
        str(a): sorted(ref_set - set(irs.filter(pl.col("anio") == a)["cvegeo"].to_list()))
        for a in (2000, 2005, 2010, 2015, 2020)
    }
    v["rezago_social"] = vi
    return v


# ---------------------------------------------------------------- main


def main() -> None:
    ap = argparse.ArgumentParser(description="CONEVAL: pobreza municipal y rezago social.")
    ap.add_argument("--descargar", action="store_true", help="vuelve a bajar los dos ZIP")
    a = ap.parse_args()
    if a.descargar:
        descargar()
    hashes = verificar_y_extraer()

    pob, est, info = leer_pobreza(
        RAW / "extraido/pobreza/Concentrado_indicadores_de_pobreza_2020.xlsx"
    )
    pob = pob.with_columns(
        pl.col("anio").cast(pl.Int16),
        pl.col("poblacion").cast(pl.Int64),
        *[pl.col(c).cast(pl.Int64) for c in pob.columns if c.endswith("_personas")],
        *[
            pl.col(c).cast(pl.Float64)
            for c in pob.columns
            if c.endswith(("_pct", "_carencias_promedio"))
        ],
    )
    pob = pob.with_columns(
        pl.lit("CONEVAL 2010-2015-2020 comparable (MCS-ENIGH 2010 / MEC 2015 / MEC 2020)").alias(
            "serie"
        )
    )
    pob = pob.sort("cvegeo", "anio")

    irs = pl.concat(
        [
            leer_irs(RAW / f"extraido/irs/IRS_entidades_mpios_{a}.xlsx", a)
            for a in (2000, 2005, 2010, 2015, 2020)
        ],
        how="vertical_relaxed",
    )
    irs = irs.with_columns(
        pl.col("anio").cast(pl.Int16),
        pl.col("poblacion_total").cast(pl.Int64),
        pl.col("lugar_nacional").cast(pl.Int32),
        pl.col("grado_rezago_social").cast(pl.Enum(GRADOS)),
        *[pl.col(c).cast(pl.Float64) for c in irs.columns if c.startswith("pct_")],
        pl.col("indice_rezago_social").cast(pl.Float64),
    ).sort("cvegeo", "anio")

    OUT.mkdir(parents=True, exist_ok=True)
    VERIF_JSON.parent.mkdir(parents=True, exist_ok=True)
    pob.write_parquet(OUT / "coneval_pobreza_municipal.parquet")
    irs.write_parquet(OUT / "coneval_rezago_social.parquet")

    v = verificar(pob, est, irs)
    v["hashes"] = hashes
    v["hashes"]["coneval_pobreza_municipal.parquet"] = sha256(
        OUT / "coneval_pobreza_municipal.parquet"
    )
    v["hashes"]["coneval_rezago_social.parquet"] = sha256(OUT / "coneval_rezago_social.parquet")
    v["info_lectura"] = info
    v["columnas_pobreza"] = {c: str(t) for c, t in pob.schema.items()}
    v["columnas_irs"] = {c: str(t) for c, t in irs.schema.items()}
    VERIF_JSON.write_text(
        json.dumps(v, ensure_ascii=False, indent=1, default=str), encoding="utf-8"
    )
    print(
        json.dumps(
            {k: v[k] for k in ("pobreza", "rezago_social")},
            ensure_ascii=False,
            indent=1,
            default=str,
        )[:20000]
    )


if __name__ == "__main__":
    main()
